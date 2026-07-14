#!/usr/bin/env python3
"""Generate SASH-VPV FYP poster document from mock FYP.docx template."""
from __future__ import annotations

import shutil
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "FYP.docx"
OUTPUT = ROOT / "SASH-VPV_FYP.docx"

TITLE = (
    "SASH-VPV: Contactless Palm Vein Recognition and Biometric Payment Platform "
    "for Pakistan — An Application of Deep Learning on NIR Palm Vein Data"
)

CONTENT = {
    0: TITLE,
    1: "Supervisor",
    2: "Dr. Benish Fida (HoD Artificial Intelligence)",
    3: "Project Team",
    4: "Syed Muhammad Huzaifa Chishty",
    5: "Shanza Rahim",
    6: "Saud Akbar",
    7: "Introduction",
    8: (
        "Pakistan has strong digital identity at onboarding through CNIC and mobile wallets, "
        "but daily authentication at checkout, office gates, and merchant counters still relies "
        "on PINs, OTPs, QR codes, and fingerprint readers that fail under dry skin, manual labor, "
        "and high-throughput public use. SASH-VPV (Secure Authentication via Subcutaneous Vascular "
        "Palm-Veins) addresses this gap with a contactless NIR palm vein recognition platform: "
        "users hover their hand over an XRTECH MagicVein Plus scanner, a custom EfficientNet-B0 + CBAM "
        "matcher extracts 512-d embeddings, and VeinPay links the verified identity to wallet payments, "
        "attendance, and enterprise access — built and trained on Pakistani palm vein data."
    ),
    9: "Methodology",
    10: (
        "Dataset Collection (SASH-VPV corpus, 2,797 NIR images, 122 subjects, 244 hand-classes) | "
        "Preprocessing (landmark detection, ROI extraction, CLAHE/Frangi/Gabor enhancement) | "
        "Model Training (EfficientNet-B0 + CBAM, ArcFace loss, 512-d L2 embeddings) | "
        "Open-Set Evaluation (5-fold cross-validation, 49 held-out classes per fold) | "
        "Architecture Comparison (ResNet, ViT, Swin, ConvNeXt vs production backbone) | "
        "System Deployment (FastAPI backend, React portal, Flutter VeinPay wallet, merchant kiosk)"
    ),
    11: "Applications",
    12: (
        "VeinPay palm payments at retail merchants — palm hover replaces PIN, OTP, and QR at checkout; "
        "merchant sets amount on kiosk, identity match triggers wallet debit with full audit trail."
    ),
    13: (
        "Offline palm attendance for BPOs, corporate offices, and co-working spaces — templates and "
        "time logs stored on-device with USB payroll export, no network or monthly subscription required."
    ),
    14: (
        "Fintech and banking step-up authentication — CNIC-bound KYC enrollment links national identity "
        "to palm template; JWT-authenticated API provides spoof-resistant biometric verification for "
        "high-value transfers, replacing SMS OTP in regulated environments."
    ),
    15: (
        "               Figure 1. 5-Fold Open-Set Cross-Validation (EER / ROC-AUC)          "
        "Figure 2. Live NIR Palm Vein Capture (XRTECH MagicVein Plus)"
    ),
}

FIGURE_SOURCES = [
    ROOT / "outputs" / "figures" / "traction_kfold_openset.png",
    ROOT / "palmpay" / "src" / "xrtech" / "output" / "test_frame.jpg",
]
FIGURE_FALLBACKS = [
    ROOT / "outputs" / "figures" / "traction_training_curve.png",
    ROOT / "palmpay" / "src" / "xrtech" / "output" / "test_frame.png",
]


def _set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def _png_to_jpg_bytes(path: Path) -> bytes:
    img = Image.open(path)
    if img.mode in ("RGBA", "P"):
        img = img.convert("RGB")
    buf = BytesIO()
    img.save(buf, format="JPEG", quality=92)
    return buf.getvalue()


def _resolve_image(path: Path, fallback: Path) -> bytes:
    src = path if path.is_file() else fallback
    if not src.is_file():
        raise FileNotFoundError(f"Missing figure: {path} and fallback {fallback}")
    if src.suffix.lower() in {".jpg", ".jpeg"}:
        return src.read_bytes()
    return _png_to_jpg_bytes(src)


def replace_paragraphs(doc: Document) -> None:
    nonempty = [p for p in doc.paragraphs if p.text.strip()]
    if len(nonempty) != len(CONTENT):
        raise RuntimeError(f"Expected {len(CONTENT)} paragraphs, found {len(nonempty)}")
    for idx, paragraph in enumerate(nonempty):
        _set_paragraph_text(paragraph, CONTENT[idx])


def replace_images(docx_path: Path) -> None:
  images = []
  for primary, fallback in zip(FIGURE_SOURCES, FIGURE_FALLBACKS):
      images.append(_resolve_image(primary, fallback))

  tmp = docx_path.with_suffix(".tmp.docx")
  shutil.copy(docx_path, tmp)
  with zipfile.ZipFile(tmp, "r") as zin:
      with zipfile.ZipFile(docx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
          for item in zin.infolist():
              data = zin.read(item.filename)
              if item.filename == "word/media/image1.jpg":
                  data = images[0]
              elif item.filename == "word/media/image2.jpg":
                  data = images[1]
              zout.writestr(item, data)
  tmp.unlink()


def main() -> None:
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    shutil.copy(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    replace_paragraphs(doc)
    doc.save(OUTPUT)
    replace_images(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
